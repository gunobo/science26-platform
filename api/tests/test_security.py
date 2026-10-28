from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.config import get_settings
from app.main import app
from app.services import validation as v
from app.services.compose import rewrite, static_compose

H = {"X-S26-Request": "1"}
UPLOAD = get_settings().upload_url


@pytest.fixture(scope="module")
def client():
    return TestClient(app)


def test_get_is_not_blocked(client):
    assert client.get("/api/health").status_code == 200


def test_unsafe_without_csrf_header_rejected(client):
    r = client.post("/api/v1/auth/logout")
    assert r.status_code == 403


@pytest.mark.parametrize("headers", [
    {"Origin": "https://science26.bssm.dev"},          # 학생 페이지 출처
    {"Sec-Fetch-Site": "same-site"},                    # 같은 사이트의 다른 출처
    {"Sec-Fetch-Site": "cross-site"},
    {"Origin": "https://evil.example"},
])
def test_cross_origin_rejected_even_with_header(client, headers):
    assert client.post("/api/v1/auth/logout", headers={**H, **headers}).status_code == 403


def test_same_origin_with_header_passes(client):
    r = client.post("/api/v1/auth/logout", headers={**H, "Origin": UPLOAD, "Sec-Fetch-Site": "same-origin"})
    assert r.status_code == 200


def test_oversized_body_rejected_before_reading(client):
    r = client.post("/api/v1/projects/me/releases", headers={**H, "Content-Length": str(500 * 1024 * 1024)}, content=b"")
    assert r.status_code == 413


def test_small_api_body_limit(client):
    r = client.post("/api/v1/auth/logout", headers=H, content=b"x" * (2 * 1024 * 1024))
    assert r.status_code == 413


def test_upload_without_session_rejected_before_body(client):
    r = client.post("/api/v1/projects/me/releases", headers=H, files={"file": ("a.zip", b"x" * 1000)}, data={"type": "static"})
    assert r.status_code == 401


def _project(tmp_path: Path, compose: dict) -> Path:
    import yaml
    (tmp_path / "docker-compose.yml").write_text(yaml.safe_dump(compose))
    return tmp_path


def test_student_labels_whitelisted(tmp_path):
    compose = {"services": {"web": {"image": "nginx", "labels": {
        "science26.public": "true", "science26.port": "8080",
        "com.docker.compose.project": "stu-9999-v1", "traefik.enable": "true", "anything": "x"}}}}
    root = _project(tmp_path, compose)
    spec = v.parse_compose((root / "docker-compose.yml").read_text(), root)
    out = rewrite(spec, root=root, student_no="1411", slug="x", version=1, seq=1, env={},
                  public_host="h", edge_network="e", memory_mb=256, stack_memory_mb=512, cpus=0.5, pids=128)
    labels = out["services"]["web"]["labels"]
    assert "com.docker.compose.project" not in labels and "anything" not in labels
    assert labels["science26.port"] == "8080"


def test_student_containers_have_log_limits(tmp_path):
    out = static_compose(student_no="1411", slug="x", version=1, seq=1, context=tmp_path,
                         public_host="h", edge_network="e", memory_mb=256, cpus=0.5, pids=128)
    assert out["services"]["web"]["logging"]["options"]["max-size"] == "5m"


def test_login_rate_limited_per_ip(client, monkeypatch):
    from app import security

    counts: dict[str, int] = {}

    async def fake_rate_limit(key, limit, window):
        counts[key] = counts.get(key, 0) + 1
        if counts[key] > limit:
            raise security.APIError(429, "RATE_LIMITED", "too many")

    monkeypatch.setattr(security, "rate_limit", fake_rate_limit)
    monkeypatch.setattr("app.routers.auth.redis.set", lambda *a, **k: _noop())

    from app.db import get_db

    class NoDB:
        async def scalar(self, *a, **k):
            return None

    async def fake_db():
        yield NoDB()

    app.dependency_overrides[get_db] = fake_db
    try:
        statuses = [client.get("/api/v1/auth/google/login", follow_redirects=False).status_code for _ in range(302)]
    finally:
        app.dependency_overrides.pop(get_db, None)
    assert statuses[0] == 307 and statuses[299] == 307 and statuses[-1] == 429


async def _noop():
    return True


def test_upload_blocked_when_disk_full(client, monkeypatch):
    import shutil
    from collections import namedtuple

    from app.routers import projects

    Usage = namedtuple("Usage", "total used free")
    monkeypatch.setattr(shutil, "disk_usage", lambda _: Usage(100, 95, 5))
    import asyncio
    with pytest.raises(projects.APIError) as e:
        asyncio.run(projects._check_upload_allowed(None, None))
    assert e.value.status == 507


def test_session_cookie_uses_host_prefix_in_production():
    from app.security import SESSION_COOKIE

    # 기본 설정(cookie_secure=True)에서는 하위 도메인이 덮어쓸 수 없는 __Host- 쿠키
    assert SESSION_COOKIE == "__Host-s26_session"


def test_traefik_dynamic_file_strips_service_worker_allowed():
    import yaml
    from pathlib import Path

    conf = yaml.safe_load((Path(__file__).parents[2] / "infra/traefik/dynamic/students.yml").read_text())
    hdr = conf["http"]["middlewares"]["student-response-headers"]["headers"]["customResponseHeaders"]
    assert hdr == {"Service-Worker-Allowed": ""}


def test_discord_mentions_only_configured_users(monkeypatch):
    import asyncio

    from app import notify

    sent = {}

    class FakeClient:
        def __init__(self, *a, **k): ...
        async def __aenter__(self): return self
        async def __aexit__(self, *a): ...
        async def post(self, url, json): sent.update(json)

    s = get_settings()
    monkeypatch.setattr(s, "discord_webhook_url", "https://discord.example/hook")
    monkeypatch.setattr(s, "discord_mention_user_ids", "945179735003127851, x")
    monkeypatch.setattr(notify.httpx, "AsyncClient", FakeClient)

    asyncio.run(notify.send("프로젝트 등록: @everyone 학생", mention=True))
    assert sent["content"].startswith("<@945179735003127851> [science26]")
    assert sent["allowed_mentions"] == {"parse": [], "users": ["945179735003127851"]}

    asyncio.run(notify.send("디스크 경고"))
    assert sent["allowed_mentions"] == {"parse": [], "users": []} and "<@" not in sent["content"]


def test_runtime_action_claimed_once(monkeypatch):
    import asyncio

    from app import queue
    from app.errors import APIError

    store: dict[str, str] = {}

    async def fake_set(key, value, nx=False, ex=None):
        if nx and key in store:
            return None
        store[key] = value
        return True

    monkeypatch.setattr(queue.redis, "set", fake_set)
    asyncio.run(queue.claim_runtime_action("1411"))
    with pytest.raises(APIError) as e:
        asyncio.run(queue.claim_runtime_action("1411"))
    assert e.value.status == 409
    asyncio.run(queue.claim_runtime_action("1203"))  # 다른 학생은 영향 없음


def test_student_caps_are_minimal(tmp_path):
    out = static_compose(student_no="1411", slug="x", version=1, seq=1, context=tmp_path,
                         public_host="h", edge_network="e", memory_mb=256, cpus=0.5, pids=128)
    web = out["services"]["web"]
    assert web["cap_drop"] == ["ALL"]
    assert sorted(web["cap_add"]) == ["CHOWN", "NET_BIND_SERVICE", "SETGID", "SETUID"]
    assert web["read_only"] is True and "no-new-privileges:true" in web["security_opt"]


def test_students_export_xlsx(client, monkeypatch):
    import io

    from openpyxl import load_workbook

    from app.db import get_db
    from app.routers import admin
    from app.security import admin_user

    def row(no, name, state, test=False):
        return {"student_no": no, "grade": no[0], "class": no[1], "number": no[2:], "name": name,
                "email": f"s{no}@bssm.hs.kr", "logged_in": state != "not_logged_in", "is_test": test, "memo": "",
                "state": state, "project": None}

    async def fake_rows(db):
        return [row("1401", "김하린", "deployed"), row("1402", "=HYPERLINK(\"x\")", "no_upload"),
                row("1403", "이서준", "not_logged_in"), row("1501", "박도윤", "no_upload"),
                row("9901", "테스트", "no_upload", test=True)]

    async def fake_db():
        yield None

    monkeypatch.setattr(admin, "_student_rows", fake_rows)
    app.dependency_overrides[admin_user] = lambda: None
    app.dependency_overrides[get_db] = fake_db
    try:
        r = client.get("/api/v1/admin/students/export.xlsx?room=1-4&filter=missing")
        bad = client.get("/api/v1/admin/students/export.xlsx?room=x")
    finally:
        app.dependency_overrides.clear()

    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    data = [row for row in ws.iter_rows(min_row=2, max_row=3, values_only=True)]
    assert [d[3] for d in data] == ["1402", "1403"]          # 1학년 4반의 미제출·미접속만, 테스트 계정 제외
    assert data[0][4] == "'=HYPERLINK(\"x\")"                 # 수식으로 실행되지 않게
    assert bad.status_code == 422


def test_project_patch_route_and_newline_normalization():
    from app.routers import projects

    routes = {(r.path, tuple(r.methods)): r.endpoint.__name__ for r in projects.router.routes}
    assert routes[("/projects/me", ("PATCH",))] == "update_project"
    assert projects._norm("가\r\n나\r다 ") == "가\n나\n다"


class _FakeRedis:
    def __init__(self):
        self.d: dict[str, str] = {}

    async def get(self, k):
        return self.d.get(k)

    async def set(self, k, v, nx=False, ex=None):
        if nx and k in self.d:
            return None
        self.d[k] = str(v)
        return True

    async def incr(self, k):
        self.d[k] = str(int(self.d.get(k, 0)) + 1)
        return int(self.d[k])

    async def expire(self, k, s):
        return True

    async def mget(self, keys):
        return [self.d.get(k) for k in keys]

    def pipeline(self):
        outer = self

        class P:
            def __init__(self):
                self.ops = []

            def incr(self, k):
                self.ops.append(("incr", k))

            def expire(self, k, s):
                pass

            async def execute(self):
                for _, k in self.ops:
                    await outer.incr(k)

        return P()


def test_visit_counts_once_per_session(monkeypatch):
    from app.db import get_db
    from app.services import visits

    fake = _FakeRedis()
    monkeypatch.setattr(visits, "redis", fake)

    async def fake_no(db, slug):
        return "1411" if slug == "orbital" else None

    monkeypatch.setattr(visits, "student_no_for", fake_no)

    async def nodb():
        yield None

    app.dependency_overrides[get_db] = nodb
    ua = {"User-Agent": "Mozilla/5.0 Chrome/130", "X-Forwarded-Uri": "/students/orbital/"}
    try:
        c1 = TestClient(app)
        r = c1.get("/api/v1/public/visit", headers=ua)
        assert r.status_code == 200 and "__Host-s26v=" in r.headers["set-cookie"]
        assert "HttpOnly" in r.headers["set-cookie"] and "Secure" in r.headers["set-cookie"]
        vid = r.cookies.get("__Host-s26v") or r.headers["set-cookie"].split("=", 1)[1].split(";")[0]
        same = {**ua, "Cookie": f"__Host-s26v={vid}"}
        c1.get("/api/v1/public/visit", headers=same)                      # 같은 세션 새로고침
        c1.get("/api/v1/public/visit", headers={**same, "X-Forwarded-Uri": "/students/orbital/level2"})
        TestClient(app).get("/api/v1/public/visit", headers=ua)           # 다른 브라우저(새 세션)
        TestClient(app).get("/api/v1/public/visit", headers={**ua, "User-Agent": "Googlebot/2.1"})
        post = TestClient(app).post("/api/v1/public/visit", headers=ua)   # 폼 POST 페이지 이동도 막지 않음
        bad = TestClient(app).get("/api/v1/public/visit", headers={**ua, "X-Forwarded-Uri": "/students/nobody/"})
    finally:
        app.dependency_overrides.pop(get_db, None)
    assert post.status_code == 200 and bad.status_code == 200
    import asyncio
    # 새 세션 3번(c1, 다른 브라우저, POST) - 로봇 제외
    assert asyncio.run(visits.counts(["1411"]))["1411"]["total"] == 3


def test_redeploy_all_skips_busy_and_blocks_double_click(monkeypatch):
    import uuid

    from app.db import get_db
    from app.routers import admin
    from app.security import admin_user

    class P:
        def __init__(self, no):
            self.student_no, self.active_release_id = no, uuid.uuid4()

    projects = [P("1201"), P("1202"), P("1203")]

    class Res:
        def scalars(self):
            return self

        def all(self):
            return projects

    class DB:
        def add(self, *_): ...
        async def execute(self, *_): return Res()
        async def commit(self): ...

    async def fake_db():
        yield DB()

    locks, jobs = {"1202"}, []

    async def fake_lock(no, owner):
        if no in locks:
            raise admin.APIError(409, "BUILD_IN_PROGRESS", "busy")
        locks.add(no)

    async def fake_enqueue(job, *args):
        jobs.append((job, args[0]))

    fake = _FakeRedis()
    monkeypatch.setattr(admin, "redis", fake)
    monkeypatch.setattr(admin, "acquire_project_lock", fake_lock)
    monkeypatch.setattr(admin, "enqueue", fake_enqueue)
    app.dependency_overrides[admin_user] = lambda: None
    app.dependency_overrides[get_db] = fake_db
    try:
        c = TestClient(app)
        r = c.post("/api/v1/admin/projects/redeploy-all", headers={**H, "Origin": UPLOAD})
        again = c.post("/api/v1/admin/projects/redeploy-all", headers={**H, "Origin": UPLOAD})
    finally:
        app.dependency_overrides.clear()
    assert r.status_code == 202 and r.json() == {"queued": ["1201", "1203"], "skipped": ["1202"]}
    assert [j for j, _ in jobs] == ["deploy_existing", "deploy_existing"]
    assert again.status_code == 409


def test_delete_student_requires_confirm_and_protects_admins(monkeypatch):
    import uuid

    from app.db import get_db
    from app.models import Role
    from app.routers import admin
    from app.security import admin_user

    me = type("U", (), {"id": uuid.uuid4(), "role": Role.admin})()
    users = {
        "1411": type("U", (), {"id": uuid.uuid4(), "role": Role.student, "is_test": False, "name": "김하린", "email": "s1411@bssm.hs.kr"})(),
        "1203": type("U", (), {"id": uuid.uuid4(), "role": Role.admin, "is_test": False, "name": "관리자", "email": "t@bssm.hs.kr"})(),
    }
    current = {"no": None}

    class DB:
        def add(self, *_): ...
        async def scalar(self, q): return users.get(current["no"])
        async def commit(self): ...

    async def fake_db():
        yield DB()

    jobs = []

    async def fake_lock(no, owner): ...
    async def fake_enqueue(job, *args): jobs.append((job, args))

    monkeypatch.setattr(admin, "acquire_project_lock", fake_lock)
    monkeypatch.setattr(admin, "enqueue", fake_enqueue)
    app.dependency_overrides[admin_user] = lambda: me
    app.dependency_overrides[get_db] = fake_db
    h = {**H, "Origin": UPLOAD}
    try:
        c = TestClient(app)
        current["no"] = "1411"
        wrong = c.delete("/api/v1/admin/users/1411?confirm=1412", headers=h)
        ok = c.delete("/api/v1/admin/users/1411?confirm=1411", headers=h)
        current["no"] = "1203"
        adm = c.delete("/api/v1/admin/users/1203?confirm=1203", headers=h)
    finally:
        app.dependency_overrides.clear()
    assert wrong.status_code == 400
    assert ok.status_code == 202 and jobs == [("delete_account", ("1411", str(users["1411"].id)))]
    assert adm.status_code == 403


def test_notice_and_memo_need_login(client):
    assert client.get("/api/v1/notice").status_code == 401
    r = client.put("/api/v1/admin/users/1411/memo", json={"memo": "x"}, headers={**H, "Origin": UPLOAD})
    assert r.status_code == 401


def test_hot_score_prefers_recent_visits():
    from app.services.visits import hot_score

    recent = hot_score([10] + [0] * 13, 10)       # 오늘 10명
    old = hot_score([0] * 10 + [10] + [0] * 3, 10)  # 10일 전 10명
    assert recent > old
    # 누적만 큰 프로젝트가 요즘 꾸준히 보는 프로젝트를 이기지 못한다
    assert hot_score([8] * 7 + [0] * 7, 56) > hot_score([0] * 14, 500)
    assert hot_score([0] * 14, 0) == 0
