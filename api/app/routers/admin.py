import csv
import io
import json
import re
import shutil
from datetime import datetime
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, Query, Response, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, select, true
from sqlalchemy.ext.asyncio import AsyncSession

from app import settings_store
from app.config import get_settings
from app.db import get_db
from app.errors import APIError, not_found, upload_invalid
from app.models import AuditLog, Project, Release, Role, User
from app.queue import acquire_project_lock, claim_runtime_action, enqueue, lock_key
from app.redis_client import redis
from app.security import admin_user, audit
from app.services import visits
from app.services.projects import project_out, student_url

router = APIRouter(prefix="/admin", tags=["admin"])


@router.get("/projects")
async def all_projects(_: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    rows = await db.execute(
        select(Project, User).join(User, User.id == Project.owner_id).order_by(Project.student_no)
    )
    rows = rows.all()
    visit_counts = await visits.counts([p.student_no for p, _ in rows])
    # 빌드·배포·재배포 잠금이 걸려 있으면 처리 중 (전체 재배포 진행률 표시용)
    locks = await redis.mget([lock_key(p.student_no) for p, _ in rows]) if rows else []
    busy = {p.student_no for (p, _), lock in zip(rows, locks) if lock}
    out = []
    for project, owner in rows:
        active = await db.get(Release, project.active_release_id) if project.active_release_id else None
        stats = await redis.get(f"stats:{project.student_no}")
        out.append({
            **project_out(project, active),
            "owner": {"email": owner.email, "name": owner.name},
            "visits": visit_counts.get(project.student_no, {"total": 0, "today": 0}),
            "busy": project.student_no in busy,
            "stats": json.loads(stats) if stats else None,
        })
    return out


@router.get("/students")
async def students(_: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    """학번이 등록된 학생 전원과 제출 현황 (프로젝트가 없는 학생도 포함).

    학번 4자리 = 학년 1자리 + 반 1자리 + 번호 2자리 (예: 1411 → 1학년 4반 11번).
    """
    return await _student_rows(db)


STATE_LABELS = {
    "deployed": "제출 완료", "building": "빌드 중", "failed": "배포 실패",
    "no_upload": "미제출", "not_logged_in": "미접속",
}


@router.get("/students/export.xlsx")
async def students_export(
    room: str = Query("all", pattern=r"^(all|test|\d-\d)$"),
    filter: str = Query("missing", pattern="^(all|missing|not_logged_in|deployed)$"),
    _: User = Depends(admin_user),
    db: AsyncSession = Depends(get_db),
):
    """학생 현황 화면과 같은 조건(학반·필터)으로 엑셀 파일을 만든다."""
    rows = [
        r for r in await _student_rows(db)
        if (r["is_test"] if room == "test" else (not r["is_test"] and (room == "all" or f"{r['grade']}-{r['class']}" == room)))
        and (
            filter == "all"
            or (filter == "missing" and r["state"] != "deployed")
            or r["state"] == filter
        )
    ]
    filter_label = {"all": "전체", "missing": "미제출", "not_logged_in": "미접속", "deployed": "제출완료"}[filter]
    room_label = {"all": "전체", "test": "테스트계정"}.get(room) or f"{room[0]}학년{room[2]}반"

    wb = Workbook()
    ws = wb.active
    ws.title = filter_label
    headers = ["학년", "반", "번호", "학번", "이름", "이메일", "상태", "로그인 여부", "프로젝트 제목", "주소", "메모"]
    ws.append(headers)
    for r in rows:
        p = r["project"] or {}
        ws.append([_cell(v) for v in (
            int(r["grade"]), int(r["class"]), int(r["number"]), r["student_no"], r["name"], r["email"],
            STATE_LABELS.get(r["state"], r["state"]), "O" if r["logged_in"] else "X", p.get("title") or "", p.get("url") or "", r["memo"],
        )])
    head_fill = PatternFill("solid", fgColor="064E3B")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDEFGHIJK", (6, 5, 6, 8, 12, 28, 10, 11, 28, 48, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    ws.append([])
    ws.append([f"{room_label} · {filter_label} {len(rows)}명 · {datetime.now(KST):%Y-%m-%d %H:%M} 기준"])

    buf = io.BytesIO()
    wb.save(buf)
    name = f"science26_{room_label}_{filter_label}_{datetime.now(KST):%Y%m%d}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"science26.xlsx\"; filename*=UTF-8''{quote(name)}"},
    )


def _cell(v):
    # 이름·제목이 = + - @ 로 시작하면 엑셀이 수식으로 실행하므로 글자로 고정한다 (CSV/수식 주입 방지)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


async def _student_rows(db: AsyncSession) -> list[dict]:
    rows = await db.execute(
        select(User, Project).outerjoin(Project, Project.owner_id == User.id)
        .where(User.student_no.is_not(None)).order_by(User.student_no)
    )
    out = []
    for user, project in rows:
        if project is None:
            state = "not_logged_in" if user.last_login_at is None else "no_upload"
        elif project.active_release_id:
            state = "deployed"
        elif project.status.value in ("building",):
            state = "building"
        else:
            state = "failed"
        out.append({
            "student_no": user.student_no,
            "grade": user.student_no[0],
            "class": user.student_no[1],
            "number": user.student_no[2:],
            "name": user.name,
            "email": user.email,
            "logged_in": user.last_login_at is not None,
            # 자동 점검(E2E) 계정은 예전 버전에서 is_test 없이 만들어졌을 수 있다
            "is_test": user.is_test or user.email.endswith("@science26.local"),
            "is_admin": user.role == Role.admin,
            "memo": user.admin_memo,
            "state": state,
            "project": None if project is None else {
                "slug": project.slug, "title": project.title, "type": project.type.value if project.type else None,
                "status": project.status.value, "url": student_url(project.slug) if project.active_release_id else None,
            },
        })
    return out


@router.get("/system")
async def system(_: User = Depends(admin_user)):
    usage = shutil.disk_usage(get_settings().data_dir)
    return {"disk": {"used": usage.used, "total": usage.total, "block_ratio": get_settings().disk_block_ratio}}


@router.post("/projects/redeploy-all", status_code=202)
async def redeploy_all(admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    """서비스 중인 프로젝트를 지금 버전 그대로 다시 띄운다 (빌드 없음).

    라우팅 라벨처럼 컨테이너를 만들 때 정해지는 설정이 바뀌었을 때 쓴다. docker restart 로는 라벨이 안 바뀐다.
    워커가 한 번에 하나씩 처리하고, 새 스택이 뜰 때까지 기존 스택이 응답해서 페이지가 끊기지 않는다.
    """
    if not await redis.set("redeploy-all", "1", nx=True, ex=300):
        raise APIError(409, "ACTION_IN_PROGRESS", "방금 전체 재배포를 시작했어요. 5분 뒤에 다시 시도하세요")
    projects = (
        await db.execute(select(Project).where(Project.active_release_id.is_not(None)).order_by(Project.student_no))
    ).scalars().all()
    queued, skipped = [], []
    for p in projects:
        try:
            await acquire_project_lock(p.student_no, str(p.active_release_id))
        except APIError:
            skipped.append(p.student_no)  # 빌드·배포 중
            continue
        await enqueue("deploy_existing", str(p.active_release_id))
        queued.append(p.student_no)
    audit(db, admin, "redeploy_all", f"{len(queued)}개 예약, {len(skipped)}개 건너뜀")
    await db.commit()
    return {"queued": queued, "skipped": skipped}


@router.post("/projects/{student_no}/stop", status_code=202)
async def force_stop(student_no: str, admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    project = (await db.execute(select(Project).where(Project.student_no == student_no))).scalar_one_or_none()
    if project is None:
        raise not_found()
    await claim_runtime_action(student_no)
    audit(db, admin, "admin_stop", student_no)
    await db.commit()
    await enqueue("stop_project", student_no, "stopped")
    return {"ok": True}


@router.delete("/projects/{student_no}", status_code=202)
async def delete_project(student_no: str, admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    project = (await db.execute(select(Project).where(Project.student_no == student_no))).scalar_one_or_none()
    if project is None:
        raise not_found()
    audit(db, admin, "delete", student_no)
    await db.commit()
    await enqueue("delete_project", student_no)
    return {"ok": True}


@router.post("/users/import")
async def import_users(
    file: UploadFile = File(...), admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)
):
    """CSV: email,학번,이름 (헤더 있어도 됨)."""
    text = (await file.read()).decode("utf-8-sig")
    created, updated, errors = 0, 0, []
    for lineno, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        if not row or not any(c.strip() for c in row):
            continue
        if lineno == 1 and "@" not in row[0]:
            continue  # 헤더
        if len(row) < 2:
            errors.append({"line": lineno, "message": "email,학번,이름 형식이 아닙니다"})
            continue
        email, student_no = row[0].strip().lower(), row[1].strip()
        name = row[2].strip() if len(row) > 2 else ""
        if "@" not in email or not re.fullmatch(r"\d{4}", student_no):
            errors.append({"line": lineno, "message": f"잘못된 값: {email}, {student_no}"})
            continue
        holder = (await db.execute(select(User).where(User.student_no == student_no))).scalar_one_or_none()
        if holder and holder.email != email:
            errors.append({"line": lineno, "message": f"학번 {student_no} 은 이미 {holder.email} 에 매핑됨"})
            continue
        user = (await db.execute(select(User).where(User.email == email))).scalar_one_or_none()
        if user is None:
            db.add(User(email=email, student_no=student_no, name=name))
            created += 1
        else:
            if user.student_no and user.student_no != student_no:
                errors.append({"line": lineno, "message": f"{email} 은 이미 학번 {user.student_no}"})
                continue
            user.student_no = student_no
            user.name = name or user.name
            updated += 1
        await db.flush()
    audit(db, admin, "import_users", f"{created}+{updated}")
    await db.commit()
    return {"created": created, "updated": updated, "errors": errors}


class StudentIn(BaseModel):
    email: str
    student_no: str
    name: str = ""
    is_test: bool = False


@router.post("/users", status_code=201)
async def add_student(body: StudentIn, admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    """학생 한 명 추가. is_test 면 학교 도메인이 아닌 구글 계정도 학생으로 로그인할 수 있다."""
    email, no = body.email.strip().lower(), body.student_no.strip()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise APIError(400, "UPLOAD_INVALID", "이메일 형식이 올바르지 않습니다")
    if not re.fullmatch(r"\d{4}", no):
        raise APIError(400, "UPLOAD_INVALID", "학번은 숫자 4자리입니다")
    holder = await db.scalar(select(User).where(User.student_no == no))
    if holder and holder.email != email:
        raise APIError(409, "CONFLICT", f"학번 {no} 은 이미 {holder.email} 에 연결돼 있어요")
    user = await db.scalar(select(User).where(User.email == email))
    if user and user.student_no and user.student_no != no:
        raise APIError(409, "CONFLICT", f"{email} 은 이미 학번 {user.student_no} 이에요")
    if user is None:
        user = User(email=email)
        db.add(user)
    user.student_no, user.name, user.is_test = no, body.name.strip() or user.name, body.is_test
    audit(db, admin, "add_test_user" if body.is_test else "add_user", f"{no} {email}")
    await db.commit()
    return {"student_no": no, "email": email, "is_test": body.is_test}


@router.delete("/users/{student_no}", status_code=202)
async def delete_student(
    student_no: str,
    confirm: str = Query(..., description="실수 방지: 지울 학번을 한 번 더"),
    admin: User = Depends(admin_user),
    db: AsyncSession = Depends(get_db),
):
    """학생 계정과 그 프로젝트(페이지·DB·파일·조회수)를 모두 지운다. 관리자 계정은 지울 수 없다."""
    if confirm != student_no:
        raise APIError(400, "CONFIRM_MISMATCH", "확인용 학번이 맞지 않아요")
    user = await db.scalar(select(User).where(User.student_no == student_no))
    if user is None:
        raise not_found()
    if user.role == Role.admin or user.id == admin.id:
        raise APIError(403, "FORBIDDEN", "관리자 계정은 지울 수 없어요")
    # 빌드·배포 중이면 기다리게 한다. 잠금은 워커가 삭제를 끝낸 뒤 푼다.
    await acquire_project_lock(student_no, "delete-account")
    audit(db, admin, "delete_test_user" if user.is_test else "delete_user", f"{student_no} {user.name} {user.email}")
    await db.commit()
    await enqueue("delete_account", student_no, str(user.id))
    return {"ok": True}


class MemoIn(BaseModel):
    memo: str


@router.put("/users/{student_no}/memo")
async def put_memo(student_no: str, body: MemoIn, _: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    """선생님만 보는 평가 메모. 학생에게는 보이지 않고, 활동 기록·알림도 남기지 않는다."""
    memo = body.memo.strip()
    if len(memo) > 2000:
        raise upload_invalid("메모는 2000자까지 쓸 수 있어요")
    user = await db.scalar(select(User).where(User.student_no == student_no))
    if user is None:
        raise not_found()
    user.admin_memo = memo
    await db.commit()
    return {"student_no": student_no, "memo": memo}


@router.get("/users")
async def list_users(_: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(User).order_by(User.student_no.nulls_last(), User.email))).scalars()
    return [
        {"email": u.email, "name": u.name, "student_no": u.student_no, "role": u.role.value}
        for u in rows
    ]


@router.get("/settings")
async def get_settings_(_: User = Depends(admin_user), db: AsyncSession = Depends(get_db)):
    return await settings_store.load(db)


@router.put("/settings")
async def put_settings(
    cfg: settings_store.PlatformConfig, admin: User = Depends(admin_user), db: AsyncSession = Depends(get_db)
):
    if not (64 <= cfg.memory_mb <= 2048 and 64 <= cfg.stack_memory_mb <= 4096
            and 0.1 <= cfg.cpus <= 4 and 32 <= cfg.pids <= 1024):
        raise upload_invalid("리소스 한도 값이 허용 범위를 벗어났습니다")
    cfg.notice = cfg.notice.strip()
    if len(cfg.notice) > 300:
        raise upload_invalid("공지는 300자까지 쓸 수 있어요")
    await settings_store.save(db, cfg)
    audit(db, admin, "settings", "")
    await db.commit()
    return cfg


# 관리자 화면 '최근 활동' 필터 → action 묶음
ACTIVITY_KINDS = {
    "login": ["first_login"],
    "project": ["register", "upload", "update_info", "rollback", "slug", "env", "stop", "start", "self_delete", "token_create", "token_revoke"],
    "admin": ["admin_stop", "delete", "import_users", "settings", "redeploy_all", "add_user", "add_test_user", "delete_test_user", "delete_user"],
}
TODAY_ACTIONS = ["first_login", "register", "upload", "update_info"]
KST = ZoneInfo("Asia/Seoul")


@router.get("/audit-logs")
async def audit_logs(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    kind: str = Query("all", pattern="^(all|login|project|admin)$"),
    _: User = Depends(admin_user),
    db: AsyncSession = Depends(get_db),
):
    cond = AuditLog.action.in_(ACTIVITY_KINDS[kind]) if kind != "all" else true()
    total = await db.scalar(select(func.count()).select_from(AuditLog).where(cond))
    rows = await db.execute(
        select(AuditLog, User).outerjoin(User, User.id == AuditLog.actor_id).where(cond)
        .order_by(AuditLog.id.desc()).offset((page - 1) * size).limit(size)
    )
    # 오늘(한국 시간) 집계는 페이지·필터와 상관없이
    start = datetime.now(KST).replace(hour=0, minute=0, second=0, microsecond=0)
    today = dict((await db.execute(
        select(AuditLog.action, func.count()).where(AuditLog.action.in_(TODAY_ACTIONS), AuditLog.at >= start)
        .group_by(AuditLog.action)
    )).all())
    return {
        "items": [
            {
                "at": log.at.isoformat(), "action": log.action, "target": log.target, "detail": log.detail,
                "actor": u.email if u else None, "actor_name": u.name if u else None,
                "actor_student_no": u.student_no if u else None,
            }
            for log, u in rows
        ],
        "total": total, "page": page, "size": size,
        "today": {a: today.get(a, 0) for a in TODAY_ACTIONS},
    }
